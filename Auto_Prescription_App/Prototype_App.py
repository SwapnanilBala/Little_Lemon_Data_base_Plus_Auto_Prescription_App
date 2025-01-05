
import os
import tkinter as tk
from datetime import datetime
from tkinter import ttk, Entry
import win32com.client

import win32api
from docx import Document
from openpyxl import load_workbook

current_time = datetime.now()
formatted_datetime = current_time.strftime("%d-%m-%Y %H:%M:%S")

# Function to save data into the Word document
def save_to_word(data):
    try:
        doc = Document(r"C:\Users\GOGOL\Documents\My Projects\Auto_Prescription_App\Auto_Prescription.docx")

        for paragraph in doc.paragraphs:
            if "Patients Name:" in paragraph.text:
                paragraph.text = f"Patients Name: {data['name']}"
            elif "Age:" in paragraph.text:
                paragraph.text = f"Age : {data['age']}"
            elif "Sex:" in paragraph.text:
                paragraph.text = f"Sex: {data['sex']}"
            elif "Address:" in paragraph.text:
                paragraph.text = f"Address: {data['address']}"
            elif "Short History with Complaints:" in paragraph.text:
                paragraph.text = f"Short History with Complaints: {data['history']}"
            elif "Important Clinical Findings:" in paragraph.text:
                paragraph.text = f"Important Clinical Findings: {data['findings']}"
            elif "Investigation Advised:" in paragraph.text:
                paragraph.text = f"Investigation Advised: {data['investigation']}"
            elif "Advice to follow:" in paragraph.text:
                paragraph.text = f"Advice to follow: {data['advice']}"
            elif "Date-Time:" in paragraph.text:
                paragraph.text = f"Date&Time: {formatted_datetime}"

        # Update Excel
        wb = load_workbook(r"C:\Users\GOGOL\Documents\My Projects\Auto_Prescription_App\Medical_Records.xlsx")
        sheet = wb.active
        new_row = [data["name"], data["age"], data["sex"], formatted_datetime, data["history"], data["findings"], data["investigation"], data["advice"]]
        sheet.append(new_row)
        
        # Save updates
        wb.save(r"C:\Users\GOGOL\Documents\My Projects\Auto_Prescription_App\Medical_Records.xlsx")
        doc.save(r"C:\Users\GOGOL\Downloads\Printable_Updated_Prescription.docx")
        
        Form_Status.delete(0, tk.END)
        Form_Status.insert(0, "Successful Submission of Details")
        print("Prescription saved successfully!")

    except Exception as e:
        print(f"Error: {e}")
        Form_Status.insert(0,"Our System ran into some problems")


def clear():
    entry_name.delete(0, tk.END)
    entry_age.delete(0, tk.END)
    entry_address.delete(0, tk.END)
    combo_sex.delete(0, tk.END)
    Form_Status.insert(0, "The patient details have been cleared")


# Function to get data from the form
def submit_form():
    Form_Status.delete(0)
    data = {
        "name": entry_name.get(),
        "age": entry_age.get(),
        "sex": combo_sex.get(),
        "address": entry_address.get(),
        "history": text_history.get("1.0", tk.END).strip(),
        "findings": text_findings.get("1.0", tk.END).strip(),
        "investigation": text_investigation.get("1.0", tk.END).strip(),
        "advice": text_advice.get("1.0", tk.END).strip()
    }
    save_to_word(data)


def print_file():
    # Check if the Word document exists
    Form_Status.delete(0)
    word_file = r"C:\Users\GOGOL\Downloads\Printable_Updated_Prescription.docx"
    if os.path.exists(word_file):
        try:
            # Launch Word and convert DOCX to PDF
            word = win32com.client.Dispatch("Word.Application")
            doc = word.Documents.Open(word_file)

            # Define the path for the PDF output
            pdf_file = r"C:\Users\GOGOL\Downloads\Printable_Updated_Prescription.pdf"

            # Save the Word document as a PDF
            doc.SaveAs(pdf_file, FileFormat=17)  # FileFormat=17 is for PDF
            doc.Close()
            word.Quit()

            # Print the PDF
            win32api.ShellExecute(0, "print", pdf_file, None, ".", 0)

            # Update form status
            Form_Status.insert(0, "The document has been converted to PDF and sent to the printer.")
        except Exception as e:
            print(f"Error: {e}")
            Form_Status.insert(0, "Error: Could not convert and print the document.")
    else:
        Form_Status.delete(0)
        Form_Status.insert(0, "Error: The Word file does not exist.")

# Tkinter GUI
root = tk.Tk()
root.title("Auto Prescription Form")
root.geometry("1600x900")
root.configure(bg="#eef1f4")  # Light grayish background color

# Title Label
title_label = tk.Label(
    root, text="Auto Prescription Form", font=("Helvetica", 22, "bold"), bg="#eef1f4", fg="#1e3d58"
)
title_label.pack(pady=20)

# Main Frame
main_frame = tk.Frame(root, bg="#eef1f4", padx=15, pady=15)
main_frame.pack(fill=tk.BOTH, expand=True)

# Patient Details Frame
details_frame = tk.LabelFrame(main_frame, text="Patient Details", font=("Helvetica", 14, "bold"), bg="#ffffff", fg="#1e3d58", padx=20, pady=20, bd=2, relief="solid")
details_frame.pack(fill=tk.X, pady=15)

# Fields for input
ttk.Label(details_frame, text="Patient's Name:").grid(row=0, column=0, padx=10, pady=10)
entry_name = ttk.Entry(details_frame, width=30, font=("Helvetica", 12))
entry_name.grid(row=0, column=1, padx=10, pady=10)

ttk.Label(details_frame, text="Age:").grid(row=1, column=0, padx=10, pady=10)
entry_age = ttk.Entry(details_frame, width=30, font=("Helvetica", 12))
entry_age.grid(row=1, column=1, padx=10, pady=10)

ttk.Label(details_frame, text="Sex:").grid(row=2, column=0, padx=10, pady=10)
combo_sex = ttk.Combobox(details_frame, values=["Male", "Female", "Other"], width=27, font=("Helvetica", 12))
combo_sex.grid(row=2, column=1, padx=10, pady=10)

ttk.Label(details_frame, text="Address:").grid(row=3, column=0, padx=10, pady=10)
entry_address = ttk.Entry(details_frame, width=30, font=("Helvetica", 12))
entry_address.grid(row=3, column=1, padx=10, pady=10)

# History and Findings Frame
history_frame = tk.LabelFrame(main_frame, text="Medical Information", font=("Helvetica", 14, "bold"), bg="#ffffff", fg="#1e3d58", padx=20, pady=20, bd=2, relief="solid")
history_frame.pack(fill=tk.X, pady=15)

ttk.Label(history_frame, text="Short History with Complaints:").grid(row=0, column=0, padx=10, pady=10)
text_history = tk.Text(history_frame, height=3, width=40, font=("Helvetica", 12))
text_history.grid(row=0, column=1, padx=10, pady=10)

ttk.Label(history_frame, text="Important Clinical Findings:").grid(row=1, column=0, padx=10, pady=10)
text_findings = tk.Text(history_frame, height=3, width=40, font=("Helvetica", 12))
text_findings.grid(row=1, column=1, padx=10, pady=10)

ttk.Label(history_frame, text="Investigation Advised:").grid(row=2, column=0, padx=10, pady=10)
text_investigation = tk.Text(history_frame, height=3, width=40, font=("Helvetica", 12))
text_investigation.grid(row=2, column=1, padx=10, pady=10)

ttk.Label(history_frame, text="Advice to follow:").grid(row=3, column=0, padx=10, pady=10)
text_advice = tk.Text(history_frame, height=3, width=40, font=("Helvetica", 12))
text_advice.grid(row=3, column=1, padx=10, pady=10)

# Form Status Entry Field
Form_Status = Entry(main_frame, font=("Arial", 12), width=50, bd=2, relief="solid")
Form_Status.pack(padx=15, pady=10)

# Buttons - Styled
button_style = {"font": ("Helvetica", 14, "bold"), "padx": 10, "pady": 10, "width": 20}


submit_button = tk.Button(
    root, text="Submit", command=submit_form, bg="#4CAF50", fg="white", **button_style
)
submit_button.pack(pady=4)


clear_button = tk.Button(
    root, text="Clear All", command=clear, bg="#f44336", fg="white", **button_style
)
clear_button.pack(pady=4)


print_button = tk.Button(
    root, text="Print Prescription", command=print_file, bg="#008CBA", fg="white", **button_style
)
print_button.pack(pady=4)


root.mainloop()
